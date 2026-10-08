from __future__ import annotations

import io
import zipfile

import docx
import pymupdf
import openpyxl
import pytest

from app.documents.context import DocText, build_llm_context
from app.documents.processor import detect_kind, identify_sections, process_bytes
from app.documents.values import best_value, find_value_candidates, format_inr, parse_amount


@pytest.mark.parametrize("text, expected", [
    ("Rs. 30,00,000", 3_000_000),
    ("₹2.4 Crore", 24_000_000),
    ("INR 18 lakh", 1_800_000),
    ("Rs 1,50,00,000/-", 15_000_000),
    ("35 Lakhs", 3_500_000),
    ("Rs. 5 Cr.", 50_000_000),
    ("₹ 22,00,000.00", 2_200_000),
    ("2 crores", 20_000_000),
])
def test_parse_indian_amounts(text, expected):
    assert parse_amount(text) == expected


def test_labelled_values_ignore_dates_and_distinguish_emd():
    text = ("Tender ID 2026_ABC_123 dated 07-10-2026. Estimated Cost: Rs. 28,50,000. "
            "EMD: Rs. 57,000. Tender fee Rs. 1,180. Clause 4.2.1 applies.")
    kinds = {c.kind: c.amount_inr for c in find_value_candidates(text)}
    assert kinds["TENDER_VALUE"] == 2_850_000
    assert kinds["EMD"] == 57_000
    assert best_value(text).amount_inr == 2_850_000


def test_unlabelled_numbers_are_not_values():
    assert best_value("Supply of 2026 licences for 30 days at 3 sites") is None


def test_format_inr():
    assert format_inr(24_000_000) == "₹2.40 Cr"
    assert format_inr(3_000_000) == "₹30 Lakh"
    assert format_inr(None) == "UNKNOWN"


# ------------------------------------------------------------------ extraction
def _docx(paras):
    d = docx.Document()
    for p in paras:
        d.add_paragraph(p)
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "SIEM licences", "4,00,00,000"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pdf(text):
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), text)
    out = doc.tobytes()
    doc.close()
    return out


def _xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Item", "Amount"])
    ws.append(["Implementation services", 2000000])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_extracts_docx_with_tables_and_sections(settings):
    res = process_bytes(_docx(["Scope of Work", "Deploy SIEM with 24x7 monitoring and log correlation for 500 sources.",
                               "Eligibility Criteria", "Average annual turnover of Rs. 5 crore in last 3 years required."]),
                        "nit.docx", settings)
    assert res.status == "EXTRACTED" and res.kind == "docx"
    assert "SIEM licences | 4,00,00,000" in res.text
    assert {"scope_of_work", "eligibility"} <= set(res.sections)


def test_extracts_pdf_text(settings):
    res = process_bytes(_pdf("Privileged Access Management solution for 2000 users"), "rfp.pdf", settings)
    assert res.status == "EXTRACTED" and "Privileged Access Management" in res.text and res.page_count == 1


def test_extracts_xlsx(settings):
    res = process_bytes(_xlsx(), "boq.xlsx", settings)
    assert res.status == "EXTRACTED" and "Implementation services | 2000000" in res.text


def test_zip_members_processed_without_touching_disk_paths(settings):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("../../etc/evil/boq.xlsx", _xlsx())
        zf.writestr("docs/rfp.pdf", _pdf("Data Loss Prevention"))
    res = process_bytes(buf.getvalue(), "bundle.zip", settings)
    assert res.status == "CONTAINER"
    assert [c.filename for c in res.children] == ["boq.xlsx", "rfp.pdf"]
    assert all(c.status == "EXTRACTED" for c in res.children)


def test_zip_bomb_limits(settings):
    settings.MAX_ZIP_TOTAL_BYTES = 1000
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("big.txt", "0" * 50_000)
    res = process_bytes(buf.getvalue(), "bomb.zip", settings)
    assert res.status == "REJECTED_UNSAFE"


def test_executables_rejected_regardless_of_extension(settings):
    res = process_bytes(b"MZ\x90\x00" + b"\x00" * 100, "tender.pdf", settings)
    assert res.status == "REJECTED_UNSAFE"


def test_legacy_doc_marked_unsupported(settings):
    res = process_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100, "old.doc", settings)
    assert res.status == "UNSUPPORTED" and "convert" in res.error


def test_detect_kind_uses_magic_bytes():
    assert detect_kind(b"%PDF-1.7", "x.docx") == "pdf"


def test_sections_identified_from_headings():
    text = "INTRODUCTION\nblah\nSCOPE OF WORK\n" + "Conduct VAPT of 20 applications. " * 3 + \
           "\nPRE-QUALIFICATION CRITERIA\n" + "Turnover above 2 crore in each of last three years. " * 2
    secs = identify_sections(text)
    assert "VAPT" in secs["scope_of_work"] and "Turnover" in secs["eligibility"]


def test_llm_context_is_targeted_when_documents_are_large():
    filler = "General conditions of contract. " * 4000
    doc = DocText("rfp.pdf", "Title page\n" + filler, {"scope_of_work": "Scope: SIEM deployment", "eligibility": "Turnover 5 Cr"})
    ctx, truncated = build_llm_context([doc], max_chars=10_000)
    assert truncated and len(ctx) <= 10_500
    assert "SIEM deployment" in ctx and "Turnover 5 Cr" in ctx


def test_llm_context_sends_small_documents_whole():
    ctx, truncated = build_llm_context([DocText("a.txt", "short tender", {})], max_chars=10_000)
    assert not truncated and "short tender" in ctx
