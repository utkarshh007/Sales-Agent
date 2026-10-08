"""Document pipeline (section 9): safety check -> type detection -> text extraction (with OCR for
scanned pages) -> normalisation -> section identification.

Documents are treated as untrusted input: type is detected from magic bytes (not the extension),
executables are rejected, archives are bounded (member count, total size, nesting) and never
extracted to disk paths taken from the archive, and macros are never executed.
"""
from __future__ import annotations

import csv
import hashlib
import io
import logging
import re
import shutil
import socket
import struct
import zipfile
from dataclasses import dataclass, field

from app.config import Settings

log = logging.getLogger(__name__)

SUPPORTED = {"pdf", "docx", "xlsx", "csv", "txt", "zip", "image"}


@dataclass
class ExtractionResult:
    filename: str
    kind: str
    status: str  # EXTRACTED | UNSUPPORTED | REJECTED_UNSAFE | FAILED | CONTAINER
    text: str = ""
    method: str | None = None
    page_count: int | None = None
    sections: dict[str, str] = field(default_factory=dict)
    error: str | None = None
    warnings: list[str] = field(default_factory=list)
    sha256: str = ""
    size: int = 0
    children: list[ExtractionResult] = field(default_factory=list)


# ------------------------------------------------------------------ type detection / safety
def detect_kind(data: bytes, filename: str) -> str:
    head = data[:8]
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"MZ") or head.startswith(b"\x7fELF") or head.startswith(b"#!"):
        return "executable"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "ole"  # legacy .doc/.xls
    if head.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                names = set(zf.namelist())
        except zipfile.BadZipFile:
            return "unknown"
        if "word/document.xml" in names:
            return "docx"
        if "xl/workbook.xml" in names:
            return "xlsx"
        return "zip"
    if head[:3] == b"\xff\xd8\xff" or head.startswith(b"\x89PNG") or head[:4] in (b"II*\x00", b"MM\x00*"):
        return "image"
    if ext == "csv":
        return "csv"
    try:
        data[:4096].decode("utf-8")
        return "txt"
    except UnicodeDecodeError:
        return "unknown"


def clamav_scan(data: bytes, settings: Settings) -> str | None:
    """Return a signature name if clamd flags the bytes, None if clean. No-op if not configured."""
    if not settings.CLAMAV_HOST:
        return None
    with socket.create_connection((settings.CLAMAV_HOST, settings.CLAMAV_PORT), timeout=60) as s:
        s.sendall(b"zINSTREAM\0")
        for i in range(0, len(data), 1 << 16):
            chunk = data[i:i + (1 << 16)]
            s.sendall(struct.pack("!L", len(chunk)) + chunk)
        s.sendall(struct.pack("!L", 0))
        reply = s.recv(4096).decode(errors="replace").strip("\0\n ")
    if reply.endswith("FOUND"):
        return reply.rsplit(":", 1)[-1].replace("FOUND", "").strip()
    if not reply.endswith("OK"):
        raise RuntimeError(f"clamd error: {reply}")
    return None


# ------------------------------------------------------------------ extraction
def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)  # de-hyphenate line breaks
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _tesseract_available() -> bool:
    return shutil.which("tesseract") is not None


def _ocr_image_bytes(img_bytes: bytes, lang: str) -> str:
    import pytesseract
    from PIL import Image
    with Image.open(io.BytesIO(img_bytes)) as im:
        return pytesseract.image_to_string(im, lang=lang)


def _extract_pdf(data: bytes, settings: Settings, res: ExtractionResult) -> None:
    import pymupdf

    ocr_ok = settings.OCR_ENABLED and _tesseract_available()
    parts, ocr_pages = [], 0
    with pymupdf.open(stream=data, filetype="pdf") as doc:
        if doc.needs_pass:
            res.status, res.error = "FAILED", "PDF is password-protected"
            return
        res.page_count = doc.page_count
        for i, page in enumerate(doc):
            t = page.get_text("text") or ""
            if len(t.strip()) < 30:  # probably a scanned page
                if ocr_ok and ocr_pages < 80:
                    pix = page.get_pixmap(dpi=200)
                    t = _ocr_image_bytes(pix.tobytes("png"), settings.OCR_LANG)
                    ocr_pages += 1
                elif not ocr_ok:
                    res.warnings.append(f"page {i + 1} looks scanned but OCR is unavailable")
            parts.append(f"\n[page {i + 1}]\n{t}")
    res.text = "".join(parts)
    res.method = "pdf_text+ocr" if ocr_pages else "pdf_text"


def _extract_docx(data: bytes, res: ExtractionResult) -> None:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    res.text, res.method = "\n".join(parts), "docx"


def _extract_xlsx(data: bytes, res: ExtractionResult) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"\n[sheet {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(cells):
                parts.append(" | ".join(cells))
    wb.close()
    res.text, res.method = "\n".join(parts), "xlsx"


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def process_bytes(data: bytes, filename: str, settings: Settings, _depth: int = 0) -> ExtractionResult:
    res = ExtractionResult(filename=filename, kind="unknown", status="FAILED",
                           sha256=hashlib.sha256(data).hexdigest(), size=len(data))
    if len(data) > settings.MAX_DOCUMENT_BYTES:
        res.status, res.error = "REJECTED_UNSAFE", f"exceeds {settings.MAX_DOCUMENT_BYTES} bytes"
        return res
    kind = detect_kind(data, filename)
    res.kind = kind
    if kind == "executable":
        res.status, res.error = "REJECTED_UNSAFE", "executable content is never processed"
        return res
    try:
        sig = clamav_scan(data, settings)
    except Exception as e:  # scanning configured but failing => do not process blindly
        res.status, res.error = "FAILED", f"malware scan failed: {e}"
        return res
    if sig:
        res.status, res.error = "REJECTED_UNSAFE", f"malware detected: {sig}"
        return res

    try:
        if kind == "pdf":
            _extract_pdf(data, settings, res)
        elif kind == "docx":
            _extract_docx(data, res)
        elif kind == "xlsx":
            _extract_xlsx(data, res)
        elif kind in ("txt", "csv"):
            text = _decode(data)
            if kind == "csv":
                rows = csv.reader(io.StringIO(text))
                text = "\n".join(" | ".join(r) for r in rows)
            res.text, res.method = text, kind
        elif kind == "image":
            if settings.OCR_ENABLED and _tesseract_available():
                res.text, res.method = _ocr_image_bytes(data, settings.OCR_LANG), "ocr"
            else:
                res.status, res.error = "UNSUPPORTED", "image document but OCR (tesseract) is unavailable"
                return res
        elif kind == "zip":
            return _process_zip(data, filename, settings, res, _depth)
        elif kind == "ole":
            res.status, res.error = "UNSUPPORTED", "legacy .doc/.xls — convert to DOCX/XLSX/PDF and re-upload"
            return res
        else:
            res.status, res.error = "UNSUPPORTED", "unrecognised file type"
            return res
    except Exception as e:
        log.exception("extraction failed for %s", filename)
        res.status, res.error = "FAILED", f"{type(e).__name__}: {e}"
        return res

    res.text = normalize_text(res.text)
    if not res.text:
        res.status, res.error = "FAILED", "no text could be extracted"
        return res
    res.status = "EXTRACTED"
    res.sections = identify_sections(res.text)
    return res


def _process_zip(data: bytes, filename: str, settings: Settings, res: ExtractionResult, depth: int) -> ExtractionResult:
    res.method = "zip"
    if depth >= 2:
        res.status, res.error = "UNSUPPORTED", "nested archives deeper than 2 levels are not processed"
        return res
    total = 0
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        infos = [i for i in zf.infolist() if not i.is_dir()]
        if len(infos) > settings.MAX_ZIP_MEMBERS:
            res.status, res.error = "REJECTED_UNSAFE", f"archive has {len(infos)} members (limit {settings.MAX_ZIP_MEMBERS})"
            return res
        for info in infos:
            total += info.file_size
            if total > settings.MAX_ZIP_TOTAL_BYTES or info.file_size > settings.MAX_DOCUMENT_BYTES:
                res.status, res.error = "REJECTED_UNSAFE", "archive expands beyond size limits (possible zip bomb)"
                return res
            # member names are used only as labels, never as filesystem paths
            label = info.filename.replace("\\", "/").split("/")[-1][:200]
            with zf.open(info) as fh:
                child = process_bytes(fh.read(settings.MAX_DOCUMENT_BYTES + 1), label, settings, depth + 1)
            res.children.append(child)
    res.status = "CONTAINER"
    return res


# ------------------------------------------------------------------ sections
SECTION_PATTERNS: dict[str, str] = {
    "scope_of_work": r"scope of (the )?(work|services|supply)|terms of reference|statement of work|\bSOW\b|job description",
    "technical_requirements": r"technical (specifications?|requirements?|compliance)|functional requirements|minimum specifications?|bill of material",
    "eligibility": r"eligibility|pre[- ]?qualification|qualification criteria|qualifying criteria|bidder.{0,10}qualification",
    "commercial": r"bill of quantit|\bBOQ\b|price (bid|schedule)|financial bid|commercial bid|schedule of rates|estimated (cost|value)",
    "emd_fees": r"earnest money|\bEMD\b|bid security|tender fee",
    "important_dates": r"(critical|important|key) dates|schedule of (bidding|tender)|bid submission (end|closing)",
    "deliverables": r"deliverables|timelines?|project schedule|implementation plan|delivery schedule",
    "support_sla": r"service level|\bSLA\b|warranty|annual maintenance|\bAMC\b|\bATS\b|support (and|&) maintenance",
    "manpower": r"manpower|resource deployment|key personnel|team composition",
    "mandatory_documents": r"documents? to be submitted|mandatory documents|checklist|technical bid documents",
}
_SECTION_RX = {k: re.compile(v, re.IGNORECASE) for k, v in SECTION_PATTERNS.items()}
MAX_SECTION_CHARS = 25_000


def identify_sections(text: str) -> dict[str, str]:
    """Split on short heading-like lines that match a known section label."""
    lines = text.split("\n")
    marks: list[tuple[int, str]] = []
    for idx, line in enumerate(lines):
        s = line.strip()
        if not s or len(s) > 140:
            continue
        for name, rx in _SECTION_RX.items():
            if rx.search(s):
                marks.append((idx, name))
                break
    sections: dict[str, str] = {}
    for n, (idx, name) in enumerate(marks):
        end = marks[n + 1][0] if n + 1 < len(marks) else len(lines)
        body = "\n".join(lines[idx:end]).strip()
        if len(body) < 40:
            continue
        prev = sections.get(name, "")
        if len(prev) < MAX_SECTION_CHARS:
            sections[name] = (prev + "\n\n" + body).strip()[:MAX_SECTION_CHARS]
    return sections


def flatten(res: ExtractionResult) -> list[ExtractionResult]:
    out = [res]
    for c in res.children:
        out.extend(flatten(c))
    return out
